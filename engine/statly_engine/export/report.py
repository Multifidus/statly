"""Full report (SPEC §10.3) as DOCX or PDF, one section per AnalysisResult:

  heading (analysis label) -> plain-language summary -> APA sentence -> table(s) ->
  assumption checks -> optional figure (PNG rendered by the frontend's WebView).

Tables and figures are numbered consecutively across the report (overriding ApaTable.number, which
is per-result). The same section model feeds both writers so DOCX and PDF stay in step.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field

from statly_engine.export import plain, run
from statly_engine.stats import apa

# Assumption statistics bounded by 1 in magnitude print without a leading zero (APA).
_BOUNDED_SYMBOLS = {"W", "r", "V", "ρ", "τ", "R²", "η²"}
_VERDICT = {"passed": "Met", "caution": "Caution", "failed": "Not met"}


def analysis_label(result: dict) -> str:
    try:
        from statly_engine.stats import registry
        return registry.get(result["analysis_id"]).label
    except Exception:  # noqa: BLE001 - unknown / removed analysis id: fall back to the id
        return result["analysis_id"]


def assumption_runs(a: dict) -> list[dict]:
    """'Normality (Pre): Shapiro-Wilk, W = .96, p < .001. Not met. <explanation>'"""
    r = apa.Rich().t(a["label"])
    scope = (a.get("applies_to") or {}).get("label")
    if scope:
        r.t(f" ({scope})")
    r.t(": ")
    bits = []
    if a.get("test_used"):
        bits.append(apa.Rich().t(a["test_used"]["label"]))
    st = a.get("statistic")
    if st is not None:
        fmt = apa.no_zero if st["symbol"] in _BOUNDED_SYMBOLS else apa.num
        sr = apa.Rich().i(st["symbol"])
        if st.get("df"):
            sr.t("(" + ", ".join(apa.df_text(d) for d in st["df"]) + ")")
        bits.append(sr.t(" = " + fmt(st["value"], 2)))
    if a.get("p") is not None:
        bits.append(apa.Rich().p(a["p"]))
    for i, b in enumerate(bits):
        if i:
            r.t(", ")
        r.extend(b)
    if bits:
        r.t(". ")
    r.t(_VERDICT.get(a.get("verdict"), str(a.get("verdict"))) + ". ")
    r.t(a.get("explanation", ""))
    return r.runs


_CORRECTION_LABEL = {"holm": "Holm", "bonferroni": "Bonferroni", "fdr_bh": "Benjamini-Hochberg"}


def adjusted_note_runs(entry: dict, family_names: dict[str, str] | None) -> list[dict]:
    """' Holm-adjusted p = .034 (family: Attitude items).' appended beside a test's APA sentence."""
    label = _CORRECTION_LABEL.get(entry.get("correction_method"), entry.get("correction_method"))
    r = apa.Rich().t(f" {label}-adjusted ").i("p").t(apa.p_relation(entry["adjusted_p"]))
    fam_name = (family_names or {}).get(entry.get("family_id") or "")
    if fam_name:
        r.t(f" (family: {fam_name})")
    r.t(".")
    return r.runs


@dataclass
class Section:
    heading: str
    summary: str
    sentence: list[dict]
    tables: list[tuple[dict, int]] = field(default_factory=list)
    assumptions: list[list[dict]] = field(default_factory=list)
    figures: list[tuple[bytes, int, str, str | None]] = field(default_factory=list)


def sections(results: list[dict], include: dict, charts: dict[str, dict],
             test_log: dict[str, dict] | None = None,
             family_names: dict[str, str] | None = None) -> list[Section]:
    out, t_no, f_no = [], 0, 0
    for res in results:
        sec = Section(analysis_label(res), res.get("plain_language_summary", ""),
                      res.get("apa_sentence") or [])
        if not include.get("sentences", True):
            sec.sentence = []
        elif test_log and sec.sentence:
            req_id = ((res.get("inputs") or {}).get("request") or {}).get("request_id")
            entry = test_log.get(req_id) if req_id else None
            if entry and entry.get("correction_method") != "none" and entry.get("adjusted_p") is not None:
                sec.sentence = list(sec.sentence) + adjusted_note_runs(entry, family_names)
        if include.get("tables", True):
            for tbl in ([res["apa_table"]] if res.get("apa_table") else []) + list(res.get("additional_tables") or []):
                t_no += 1
                sec.tables.append((tbl, t_no))
        if include.get("assumptions", True):
            sec.assumptions = [assumption_runs(a) for a in res.get("assumptions") or []]
        if include.get("charts", True):
            req_id = ((res.get("inputs") or {}).get("request") or {}).get("request_id")
            chart = charts.get(req_id) if req_id else None
            if chart is not None:
                f_no += 1
                title = chart.get("title") or (res["apa_table"]["title"] if res.get("apa_table") else sec.heading)
                sec.figures.append((chart["png"], f_no, title, chart.get("note")))
        out.append(sec)
    return out


def _byline(author: str | None) -> str:
    d = _dt.date.today()
    date = f"{d.strftime('%B')} {d.day}, {d.year}"
    return f"{author} · {date}" if author else date


def write_docx(path: str, title: str, author: str | None, secs: list[Section]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from statly_engine.export import docx as dx

    doc = dx.new_document()
    doc.core_properties.title = title
    doc.core_properties.author = author or ""
    doc.add_paragraph(title, style="Title")
    dx.paragraph(doc, [run(_byline(author))], space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    for sec in secs:
        h = doc.add_paragraph(sec.heading, style="Heading 1")
        h.paragraph_format.space_before = Pt(12)
        if sec.summary:
            dx.paragraph(doc, [run(sec.summary)])
        if sec.sentence:
            dx.paragraph(doc, sec.sentence)
        for tbl, n in sec.tables:
            dx.add_apa_table(doc, tbl, n)
        if sec.assumptions:
            doc.add_paragraph("Assumption Checks", style="Heading 2")
            for a in sec.assumptions:
                dx.paragraph(doc, a, style="List Bullet")
        for png, n, ftitle, note in sec.figures:
            dx.add_figure(doc, png, n, ftitle, note)
    doc.save(path)


def write_pdf(path: str, title: str, author: str | None, secs: list[Section]) -> None:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import Spacer

    from statly_engine.export import pdf

    width = pdf.text_width()
    story = [pdf.para([run(title, bold=True)], pdf.style(14, TA_CENTER, space_after=6, bold=True)),
             pdf.para([run(_byline(author))], pdf.style(12, TA_CENTER, space_after=18))]
    for sec in secs:
        story.append(pdf.para([run(sec.heading, bold=True)], pdf.style(12, TA_CENTER, space_after=6, bold=True)))
        if sec.summary:
            story.append(pdf.para([run(sec.summary)], pdf.style(12, space_after=6)))
        if sec.sentence:
            story.append(pdf.para(sec.sentence, pdf.style(12, space_after=12)))
        for tbl, n in sec.tables:
            story += pdf.apa_table(tbl, n, width)
        if sec.assumptions:
            story.append(pdf.para([run("Assumption Checks", bold=True)], pdf.style(12, space_after=6, bold=True)))
            for a in sec.assumptions:
                st = pdf.style(12, space_after=4)
                st.leftIndent, st.bulletIndent = 18, 6
                story.append(pdf.Paragraph(pdf.runs_markup(a), st, bulletText="•"))
            story.append(Spacer(1, 12))
        for png, n, ftitle, note in sec.figures:
            story += pdf.figure(png, n, ftitle, note, width)
    pdf.build(path, story, title, author)


def plain_sentence(res: dict) -> str:
    return plain(res.get("apa_sentence"))
