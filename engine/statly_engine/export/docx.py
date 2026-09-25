"""python-docx rendering helpers: an APA 7 document (Times New Roman 12pt, 1-inch margins), a real
Word table style "APA Table" (top and bottom rules, no vertical or inner lines, header row repeats),
APA tables from the ApaTable contract, APA figures, and plain grids for the codebook / Test Log.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt

from statly_engine.export import (align_of, cell_runs, group_header, note_paragraphs, run, table_label)

FONT = "Times New Roman"
APA_TABLE_STYLE = "APA Table"
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
          "right": WD_ALIGN_PARAGRAPH.RIGHT}

_STYLE_XML = f"""
<w:style {nsdecls('w')} w:type="table" w:customStyle="1" w:styleId="APATable">
  <w:name w:val="{APA_TABLE_STYLE}"/>
  <w:basedOn w:val="TableNormal"/>
  <w:uiPriority w:val="99"/>
  <w:qFormat/>
  <w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>
  <w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}" w:eastAsia="{FONT}"/><w:sz w:val="24"/></w:rPr>
  <w:tblPr>
    <w:tblBorders>
      <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
      <w:left w:val="nil"/>
      <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
      <w:right w:val="nil"/>
      <w:insideH w:val="nil"/>
      <w:insideV w:val="nil"/>
    </w:tblBorders>
    <w:tblCellMar><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tblCellMar>
  </w:tblPr>
  <w:tblStylePr w:type="firstRow">
    <w:tcPr><w:tcBorders><w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/></w:tcBorders></w:tcPr>
  </w:tblStylePr>
</w:style>
"""


def new_document(landscape: bool = False):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    for name in ("Title", "Heading 1", "Heading 2", "Caption"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.color.rgb = None
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(12) if name != "Title" else Pt(14)
        st.font.bold = name != "Caption"
    title = doc.styles["Title"]
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for bdr in title.element.xpath("./w:pPr/w:pBdr"):  # default template underlines titles in blue
        bdr.getparent().remove(bdr)
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles.element.append(parse_xml(_STYLE_XML))
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    return doc


def text_width(doc) -> Emu:
    sec = doc.sections[-1]
    return Emu(sec.page_width - sec.left_margin - sec.right_margin)


def add_runs(paragraph, runs: list[dict]):
    for r in runs:
        rr = paragraph.add_run(r.get("text", ""))
        f = rr.font
        if r.get("italic"):
            f.italic = True
        if r.get("bold"):
            f.bold = True
        if r.get("subscript"):
            f.subscript = True
        if r.get("superscript"):
            f.superscript = True
    return paragraph


def paragraph(doc, runs: list[dict], style: str | None = None, space_after: int = 6, align=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, runs)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def _set_cell_border(cell, **edges):
    """edges: top/bottom=True -> single 1pt black; False -> nil."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = parse_xml(f"<w:tcBorders {nsdecls('w')}/>")
        tc_pr.append(borders)
    for edge, on in edges.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = parse_xml(f"<w:{edge} {nsdecls('w')}/>")
            borders.append(el)
        if on:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "8")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="true"/>'))


def _fill(cell, runs: list[dict], align: str, indent: int = 0, size: int | None = None):
    p = cell.paragraphs[0]
    add_runs(p, runs)
    p.alignment = _ALIGN[align]
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.left_indent = Pt(12 * indent)
    if size:
        for r in p.runs:
            r.font.size = Pt(size)


def _style_table(table):
    """Apply the APA Table style and repeat its margins on the table itself (some renderers ignore
    table-style cell margins)."""
    table.style = table.part.document.styles[APA_TABLE_STYLE]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table._tbl.tblPr.append(parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:left w:w="100" w:type="dxa"/><w:right w:w="100" w:type="dxa"/>'
        "</w:tblCellMar>"))


def _outer_rules(table):
    """Explicit top/bottom rules on the first and last rows (mirrors the style for renderers that
    ignore table-style borders)."""
    for c in table.rows[0].cells:
        _set_cell_border(c, top=True)
    for c in table.rows[-1].cells:
        _set_cell_border(c, bottom=True)


def add_apa_table(doc, table: dict, number: int | None = None, font_size: int | None = None):
    """Table N (bold) / title (italic) / table with APA rules / notes."""
    lab = paragraph(doc, [run(table_label(table, number), bold=True)], space_after=6)
    lab.paragraph_format.keep_with_next = True
    tit = paragraph(doc, [run(table.get("title", ""), italic=True)], space_after=6)
    tit.paragraph_format.keep_with_next = True
    cols = table["columns"]
    aligns = [align_of(c) for c in cols]
    groups = group_header(table)
    body = table.get("rows") or []
    n_head = 2 if groups else 1
    t = doc.add_table(rows=n_head + max(len(body), 1), cols=len(cols))
    _style_table(t)
    size = font_size or (10 if len(cols) > 8 else None)
    if groups:
        for runs, first, span in groups:
            c = t.cell(0, first)
            if span > 1:
                c = c.merge(t.cell(0, first + span - 1))
            _fill(c, runs, "center", size=size)
            if runs:
                _set_cell_border(c, bottom=True)
        _repeat_header(t.rows[0])
    hdr = t.rows[n_head - 1]
    for i, col in enumerate(cols):
        _fill(hdr.cells[i], col["header"], "left" if aligns[i] == "left" and i == 0 else "center", size=size)
        _set_cell_border(hdr.cells[i], bottom=True)
    _repeat_header(hdr)
    for ri, row in enumerate(body):
        cells = t.rows[n_head + ri].cells
        for ci, cell in enumerate(row["cells"][:len(cols)]):
            runs = cell_runs(cell)
            if row.get("kind") == "section_header" and ci == 0:
                runs = [{**r, "italic": True} for r in runs]
            _fill(cells[ci], runs, aligns[ci], row.get("indent", 0) if ci == 0 else 0, size=size)
    _outer_rules(t)
    for para in note_paragraphs(table):
        p = paragraph(doc, para, space_after=0)
        p.paragraph_format.space_before = Pt(6)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def add_grid(doc, headers: list[str], rows: list[list[list[dict] | str]], font_size: int = 10,
             widths: list[float] | None = None):
    """Plain APA-ruled grid (codebook / Test Log). Cells are strings or RichText runs."""
    t = doc.add_table(rows=1 + max(len(rows), 1), cols=len(headers))
    _style_table(t)
    for i, h in enumerate(headers):
        _fill(t.rows[0].cells[i], [run(h)], "left", size=font_size)
    _repeat_header(t.rows[0])
    for ri, r in enumerate(rows):
        for ci, v in enumerate(r):
            _fill(t.rows[ri + 1].cells[ci], [run(v)] if isinstance(v, str) else v, "left", size=font_size)
    for c in t.rows[0].cells:
        _set_cell_border(c, bottom=True)
    _outer_rules(t)
    if widths:
        t.autofit = False
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Inches(w)
    return t


def add_figure(doc, png: bytes, number: int, title: str, note: str | None = None):
    """APA 7 figure: 'Figure N' (bold), title (italic), image, optional 'Note.'."""
    from PIL import Image

    lab = paragraph(doc, [run(f"Figure {number}", bold=True)], space_after=6)
    lab.paragraph_format.keep_with_next = True
    tit = paragraph(doc, [run(title, italic=True)], space_after=6)
    tit.paragraph_format.keep_with_next = True
    with Image.open(io.BytesIO(png)) as im:
        w_px, h_px = im.size
        dpi = (im.info.get("dpi") or (96, 96))[0] or 96
    max_w = text_width(doc)
    width = min(Emu(int(w_px / float(dpi) * 914400)), max_w)
    max_h = Inches(7.5)
    if w_px and width * h_px / w_px > max_h:
        width = Emu(int(max_h * w_px / h_px))
    doc.add_picture(io.BytesIO(png), width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if note:
        paragraph(doc, [run("Note.", italic=True), run(" " + note)], space_after=6)
    doc.add_paragraph()
