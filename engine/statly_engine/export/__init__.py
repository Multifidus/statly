"""Phase 8 exports (SPEC §10.3): APA tables as clipboard HTML, report DOCX/PDF, dataset XLSX/CSV,
codebook XLSX/DOCX, Test Log XLSX/CSV/DOCX.

Everything here renders strings the engine already formatted (stats/apa.py display strings and
RichText runs); where a contract carries only raw numbers (assumption checks, Test Log summaries)
the numbers are formatted with stats/apa.py, never ad hoc. All writers are offline and write
atomically (<path>.tmp -> rename) to a user-chosen absolute path; an existing file is replaced only
when the caller passes overwrite=true.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Callable

from statly_engine.errors import FileUnreadable, InvalidParams

EXTENSIONS = {"docx": ".docx", "pdf": ".pdf", "xlsx": ".xlsx", "csv": ".csv"}


def check_target(path: object, fmt: str, overwrite: bool) -> Path:
    """Validate a save-dialog path: absolute, extension matching `fmt`, existing parent folder,
    not a directory, and no silent overwrite."""
    if not isinstance(path, str) or not path.strip():
        raise InvalidParams("Choose where to save the file.", reason="missing_path")
    target = Path(path)
    if not target.is_absolute():
        raise InvalidParams("The save location must be a full (absolute) path chosen in the save dialog.",
                            reason="relative_path", path=path)
    if ".." in target.parts:
        raise InvalidParams("The save location can't contain '..'.", reason="relative_path", path=path)
    ext = EXTENSIONS[fmt]
    if target.suffix.lower() != ext:
        raise InvalidParams(f"A {fmt.upper()} export must be saved with the '{ext}' extension.",
                            reason="wrong_extension", path=path)
    if not target.parent.is_dir():
        raise InvalidParams("The folder you chose doesn't exist.", reason="missing_folder", path=path)
    if target.is_dir():
        raise InvalidParams("That path is a folder, not a file.", reason="is_directory", path=path)
    if target.exists() and not overwrite:
        raise InvalidParams(f"'{target.name}' already exists. Confirm replacing it to overwrite.",
                            reason="file_exists", path=path)
    return target


def atomic_write(target: Path, writer: Callable[[str], None]) -> int:
    """Run writer(tmp_path), fsync, rename over target. Returns the final size in bytes."""
    tmp = target.with_name(target.name + ".tmp")
    try:
        writer(str(tmp))
        with open(tmp, "rb+") as fh:
            os.fsync(fh.fileno())
        os.replace(tmp, target)
    except OSError as exc:
        tmp.unlink(missing_ok=True)
        raise FileUnreadable(f"Could not write {target.name}: {exc.strerror or exc}.", path=str(target)) from exc
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    return target.stat().st_size


def plain(runs: list[dict] | None) -> str:
    return "".join(r.get("text", "") for r in runs or [])


def run(text: str, italic: bool = False, bold: bool = False) -> dict:
    return {"text": text, "italic": italic, "subscript": False, "superscript": False, "bold": bold}


def cell_runs(cell: dict) -> list[dict]:
    """A TableCell's styled runs: engine display strings verbatim, text cells as their runs."""
    kind = cell.get("type")
    if kind in ("number", "p_value", "interval"):
        return [run(cell["display"])]
    if kind == "text":
        return list(cell.get("text") or [])
    return []


def table_number(table: dict, number: int | None) -> int | None:
    return number if number is not None else table.get("number")


def table_label(table: dict, number: int | None) -> str:
    n = table_number(table, number)
    return f"Table {n}" if n is not None else "Table"


def note_paragraphs(table: dict) -> list[list[dict]]:
    """APA 7 notes: 'Note.' (italic) + general note; then specific notes; then probability notes."""
    notes = table.get("notes") or {}
    out: list[list[dict]] = []
    if notes.get("general"):
        out.append([run("Note.", italic=True), run(" ")] + list(notes["general"]))
    specific = [n for n in notes.get("specific") or [] if n]
    if specific:
        para: list[dict] = []
        for i, n in enumerate(specific):
            para += ([run(" ")] if i else []) + list(n)
        out.append(para)
    probability = [n for n in notes.get("probability") or [] if n]
    if probability:
        para = []
        for i, n in enumerate(probability):
            para += ([run(" ")] if i else []) + list(n)
        out.append(para)
    return out


def group_header(table: dict) -> list[tuple[list[dict], int, int]] | None:
    """Spanning header row as (runs, first_column, span) segments covering every column, or None."""
    groups = sorted(table.get("column_groups") or [], key=lambda g: g["first_column"])
    if not groups:
        return None
    ncol = len(table["columns"])
    segs, col = [], 0
    for g in groups:
        first, span = g["first_column"], g["span"]
        if first < col or first + span > ncol:
            continue  # overlapping / out of range: ignore rather than corrupt the grid
        while col < first:
            segs.append(([], col, 1))
            col += 1
        segs.append((list(g["label"]), first, span))
        col = first + span
    while col < ncol:
        segs.append(([], col, 1))
        col += 1
    return segs


def align_of(column: dict) -> str:
    """APA tables center numeric columns; 'decimal' renders centered (no decimal tabs in HTML/PDF)."""
    return {"left": "left", "right": "right"}.get(column.get("align"), "center")


def xlsx_value(ws, x, font=None, wrap: bool = False):
    """A write-only-sheet value. Strings stay text: control characters Excel rejects are removed and
    a leading '=' is stored as text rather than becoming a formula."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment

    if isinstance(x, str):
        x = ILLEGAL_CHARACTERS_RE.sub("", x)
    if not (font or wrap or (isinstance(x, str) and x.startswith("="))):
        return x
    c = WriteOnlyCell(ws, value=x)
    if isinstance(x, str) and x.startswith("="):
        c.data_type = "s"
    if font is not None:
        c.font = font
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def write_xlsx_sheets(path: str, sheets: list[tuple[str, list[str], list[list]]]) -> None:
    """Simple multi-sheet workbook: bold header row, frozen, wrapped text, sized columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            longest = max([len(h)] + [len(str(r[i - 1])) if r[i - 1] is not None else 0 for r in rows])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
        ws.append([xlsx_value(ws, h, font=bold) for h in headers])
        for r in rows:
            ws.append([xlsx_value(ws, v, wrap=isinstance(v, str) and len(v) > 40) for v in r])
    wb.save(path)
