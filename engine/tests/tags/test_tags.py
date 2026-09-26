"""Phase 9 qualitative coding (SPEC §11.1): codebook, tagging, search spans, summaries,
yes/no variables, persistence, and exports on messy_qualtrics Q10."""

from __future__ import annotations

import json
import zipfile

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook

from statly_engine.contracts import ProjectFile, TagCodebook
from statly_engine.data import project as proj
from statly_engine.data import tags
from statly_engine.data.store import ROW_ID, DatasetStore
from statly_engine.errors import InvalidParams, StaleOrUnknown
from statly_engine.rpc_methods import session_handlers

from data._helpers import MESSY, PRACTICE, import_single

EXAMPLE = PRACTICE.parents[1] / "contracts" / "examples" / "ProjectFile.json"


def frontend_project(meta: dict | None) -> dict:
    p = json.loads(EXAMPLE.read_text())
    p.pop("$schema", None)
    p.update(project_id="8f0c2a4e-1b7d-4c1e-9a55-2f6f3c1d9e02", name="Coded survey", dataset_meta=meta,
             data_path="data/dataset.parquet" if meta else None,
             test_log=[], test_families=[], chart_specs=[], tag_codebook=None, study_plan=None)
    return p

H = session_handlers()


def call(store, method, **params):
    return H[method](store, params)


@pytest.fixture
def ds(store):
    meta = import_single(store, MESSY / "messy_3header.csv")
    return meta


def _tag(store, d, name, **kw):
    return call(store, "tags.codebook.upsert", dataset_id=d, tag={"name": name, **kw})["tag"]


def test_codebook_crud(store, ds):
    d = ds["dataset_id"]
    assert call(store, "tags.codebook.get", dataset_id=d)["codebook"] == tags.empty_codebook()
    a = _tag(store, d, "Pacing", definition="Talks about speed of the course")
    b = _tag(store, d, "Recommends")
    assert a["id"] == "tag_pacing" and a["color"] == "#0072B2" and b["color"] == "#E69F00"
    with pytest.raises(InvalidParams):
        _tag(store, d, "pacing")  # names are unique, ignoring case
    with pytest.raises(InvalidParams):
        call(store, "tags.codebook.upsert", dataset_id=d, tag={"name": "Bad", "color": "blue"})
    edited = call(store, "tags.codebook.upsert", dataset_id=d,
                  tag={"id": a["id"], "name": "Pace", "color": "#009E73", "definition": "Speed"})
    assert edited["tag"] == {"id": "tag_pacing", "name": "Pace", "color": "#009E73", "definition": "Speed"}
    call(store, "tags.apply", dataset_id=d, row_id=3, variable="Q10", tag_ids=[a["id"], b["id"]])
    book = call(store, "tags.codebook.delete", dataset_id=d, tag_id=a["id"])["codebook"]
    TagCodebook.model_validate(book)
    assert [t["id"] for t in book["tags"]] == ["tag_recommends"]
    assert book["applications"] == [{"row_id": 3, "variable": "Q10", "tag_ids": ["tag_recommends"]}]
    with pytest.raises(StaleOrUnknown):
        call(store, "tags.codebook.get", dataset_id="nope")


def test_apply_and_unapply(store, ds):
    d = ds["dataset_id"]
    a, b = _tag(store, d, "A"), _tag(store, d, "B")
    res = call(store, "tags.apply", dataset_id=d, row_id=5, variable="Q10", tag_ids=[b["id"], a["id"], b["id"]])
    assert res["tag_ids"] == [a["id"], b["id"]]  # deduplicated, codebook order
    call(store, "tags.apply", dataset_id=d, row_id=5, variable="Q10", tag_ids=[b["id"]])
    book = tags.codebook(store, d)
    assert book["applications"] == [{"row_id": 5, "variable": "Q10", "tag_ids": [b["id"]]}]
    call(store, "tags.apply", dataset_id=d, row_id=5, variable="Q10", tag_ids=[])
    assert book["applications"] == []
    with pytest.raises(InvalidParams):
        call(store, "tags.apply", dataset_id=d, row_id=5, variable="Q10", tag_ids=["tag_missing"])
    with pytest.raises(InvalidParams):
        call(store, "tags.apply", dataset_id=d, row_id=5, variable="Q6", tag_ids=[a["id"]])  # not text
    with pytest.raises(StaleOrUnknown):
        call(store, "tags.apply", dataset_id=d, row_id=99999, variable="Q10", tag_ids=[a["id"]])


def test_paged_search_with_spans(store, ds):
    d = ds["dataset_id"]
    df = store.get(d).df
    text = df["Q10"].iloc[0]
    page1 = call(store, "tags.responses", dataset_id=d, variable="Q10", search='PACING "extra practice"',
                 context_variables=["Q6"], offset=0, limit=50)
    assert page1["total"] == page1["total_responses"] == len(df) and len(page1["items"]) == 50
    first = page1["items"][0]
    assert first["row_id"] == int(df[ROW_ID].iloc[0]) and first["context"]["Q6"] == int(df["Q6"].iloc[0])
    assert [text[s:e] for s, e in first["matches"]] == ["pacing", '"extra practice"'[1:-1]]
    page3 = call(store, "tags.responses", dataset_id=d, variable="Q10", search="pacing", offset=100, limit=50)
    assert [i["row_id"] for i in page3["items"]] == [int(x) for x in df[ROW_ID].iloc[100:]]
    none = call(store, "tags.responses", dataset_id=d, variable="Q10", search="pacing zebra")
    assert none["total"] == 0 and none["items"] == []
    # Q9 (distinct answers): a search narrows the list.
    q9 = call(store, "tags.responses", dataset_id=d, variable="Q9", search="student73@")
    assert q9["total"] == int(df["Q9"].str.contains("student73@", regex=False).sum()) == 1


def test_filters_and_tag_filter(store, ds):
    d = ds["dataset_id"]
    df = store.get(d).df
    a = _tag(store, d, "A")
    rows_q6_4 = df.loc[df["Q6"] == 4, ROW_ID].tolist()
    call(store, "tags.apply", dataset_id=d, row_id=int(rows_q6_4[0]), variable="Q10", tag_ids=[a["id"]])
    f = [{"variable": "Q6", "values": [4, 5]}]
    res = call(store, "tags.responses", dataset_id=d, variable="Q10", filters=f)
    assert res["total"] == int(df["Q6"].isin([4, 5]).sum())
    tagged = call(store, "tags.responses", dataset_id=d, variable="Q10", filters=f, tag_filter=a["id"])
    assert [i["row_id"] for i in tagged["items"]] == [int(rows_q6_4[0])]
    assert tagged["items"][0]["tag_ids"] == [a["id"]]
    untagged = call(store, "tags.responses", dataset_id=d, variable="Q10", filters=f, tag_filter="untagged")
    assert untagged["total"] == res["total"] - 1


def test_match_spans_are_utf16_offsets():
    text = "😀 Pacing; pacing — naïve Pacing"
    spans = tags.match_spans(text, ["pacing"])
    utf16 = text.encode("utf-16-le")
    assert [utf16[2 * s:2 * e].decode("utf-16-le") for s, e in spans] == ["Pacing", "pacing", "Pacing"]
    assert spans[0] == [3, 9]  # the emoji is two UTF-16 code units
    assert tags.match_spans("abcabc", ["abc", "bca"]) == [[0, 6]]  # overlapping matches merge
    assert tags.match_spans("abc", []) == []


def test_summary_by_group_matches_pandas(store, ds):
    d = ds["dataset_id"]
    df = store.get(d).df
    a, b = _tag(store, d, "Pacing"), _tag(store, d, "Recommends")
    rows = df[ROW_ID].tolist()
    tag_a, tag_b = set(rows[0:40:3]), set(rows[5:90:4])
    for r in tag_a | tag_b:
        ids = [t for t, s in ((a["id"], tag_a), (b["id"], tag_b)) if r in s]
        call(store, "tags.apply", dataset_id=d, row_id=int(r), variable="Q10", tag_ids=ids)
    res = call(store, "tags.summary", dataset_id=d, variable="Q10", by="Q6")

    ref = df[[ROW_ID, "Q6"]].copy()
    ref["A"] = ref[ROW_ID].isin(tag_a)
    ref["B"] = ref[ROW_ID].isin(tag_b)
    n = len(ref)
    assert res["n_responses"] == n and res["n_coded"] == len(tag_a | tag_b)
    assert [c["count"] for c in res["overall"]] == [int(ref["A"].sum()), int(ref["B"].sum())]
    assert res["overall"][0]["percent"] == round(100 * ref["A"].sum() / n, 1)
    valid = ref[ref["Q6"].notna() & (ref["Q6"] != -99)]
    grp = valid.groupby("Q6")[["A", "B"]].agg(["sum", "count"])
    assert [g["value"] for g in res["groups"]] == [int(x) for x in grp.index]
    for g in res["groups"]:
        row = grp.loc[g["value"]]
        assert g["n_responses"] == int(row[("A", "count")])
        assert [c["count"] for c in g["counts"]] == [int(row[("A", "sum")]), int(row[("B", "sum")])]
        assert g["counts"][1]["percent"] == round(100 * row[("B", "sum")] / row[("B", "count")], 1)
    assert res["n_missing_group"] == n - len(valid)


def test_to_variables_snapshot_roundtrip(store, ds, tmp_path):
    d = ds["dataset_id"]
    a, b = _tag(store, d, "Pacing"), _tag(store, d, "Peer help")
    for r in (0, 1, 2):
        call(store, "tags.apply", dataset_id=d, row_id=r, variable="Q10", tag_ids=[a["id"]])
    call(store, "tags.apply", dataset_id=d, row_id=7, variable="Q10", tag_ids=[a["id"], b["id"]])
    before = ds["snapshot_id"]
    res = call(store, "tags.to_variables", dataset_id=d, snapshot_id=before, variable="Q10")
    meta = res["dataset_meta"]
    assert meta["snapshot_id"] != before
    assert [c["variable"] for c in res["created"]] == ["Q10_pacing", "Q10_peer_help"]
    assert (res["created"][0]["n_yes"], res["created"][0]["n_no"]) == (4, len(store.get(d).df) - 4)
    v = next(x for x in meta["variables"] if x["name"] == "Q10_pacing")
    assert v["level"] == "nominal" and v["dtype"] == "integer" and [x["label"] for x in v["value_labels"]] == ["No", "Yes"]
    df = store.get(d).df
    assert df.loc[df[ROW_ID].isin([0, 1, 2, 7]), "Q10_pacing"].tolist() == [1, 1, 1, 1]
    assert int(df["Q10_peer_help"].sum()) == 1
    # stale snapshot is refused
    with pytest.raises(StaleOrUnknown):
        call(store, "tags.to_variables", dataset_id=d, snapshot_id=before, variable="Q10")
    # re-running after more tagging updates the same variables in place
    call(store, "tags.apply", dataset_id=d, row_id=9, variable="Q10", tag_ids=[b["id"]])
    again = call(store, "tags.to_variables", dataset_id=d, snapshot_id=meta["snapshot_id"], variable="Q10",
                 tag_ids=[b["id"]])
    assert again["created"] == [{"tag_id": b["id"], "variable": "Q10_peer_help", "n_yes": 2,
                                 "n_no": len(df) - 2, "n_missing": 0, "updated": True}]
    assert sum(1 for x in again["dataset_meta"]["variables"] if x["name"].startswith("Q10_peer")) == 1
    # undo goes back to the pre-variable snapshot
    restored = store.restore(d, before)
    assert "Q10_pacing" not in {x["name"] for x in restored["variables"]}
    store.restore(d, again["dataset_meta"]["snapshot_id"])
    # save -> load keeps the new variables and their values
    path = tmp_path / "tags.statly"
    proj.save(store, str(path), frontend_project(store.get(d).meta), "0.1.0")
    fresh = DatasetStore()
    loaded = proj.load(fresh, str(path))
    assert loaded["project"]["dataset_meta"]["snapshot_id"] == again["dataset_meta"]["snapshot_id"]
    pd.testing.assert_frame_equal(fresh.get(d).df, store.get(d).df, check_exact=True)


def test_project_save_load_persists_codebook(store, ds, tmp_path):
    d = ds["dataset_id"]
    a = _tag(store, d, "Pacing", definition="Speed")
    call(store, "tags.apply", dataset_id=d, row_id=4, variable="Q10", tag_ids=[a["id"]])
    path = tmp_path / "coded.statly"
    res = proj.save(store, str(path), frontend_project(ds), "0.1.0")
    ProjectFile.model_validate(res["project"])
    assert res["project"]["tag_codebook"] == tags.codebook(store, d)
    with zipfile.ZipFile(path) as zf:
        saved = json.loads(zf.read("project.json"))["tag_codebook"]
    assert saved["applications"] == [{"row_id": 4, "variable": "Q10", "tag_ids": ["tag_pacing"]}]

    fresh = DatasetStore()
    loaded = proj.load(fresh, str(path))
    assert loaded["project"]["tag_codebook"] == saved
    assert call(fresh, "tags.codebook.get", dataset_id=d)["codebook"] == saved
    item = call(fresh, "tags.responses", dataset_id=d, variable="Q10", tag_filter="tag_pacing")["items"]
    assert [i["row_id"] for i in item] == [4]
    # autosave carries the codebook too
    call(fresh, "tags.apply", dataset_id=d, row_id=6, variable="Q10", tag_ids=[a["id"]])
    auto = proj.autosave(fresh, str(tmp_path / "auto"), str(path), loaded["project"], "0.1.0")
    with zipfile.ZipFile(auto["autosave_path"]) as zf:
        assert len(json.loads(zf.read("project.json"))["tag_codebook"]["applications"]) == 2
    # a project without tags keeps tag_codebook null
    other = DatasetStore()
    meta2 = import_single(other, MESSY / "messy_3header.csv")
    out = proj.save(other, str(tmp_path / "plain.statly"), frontend_project(meta2), "0.1.0")
    assert out["project"]["tag_codebook"] is None


def test_exports(store, ds, tmp_path):
    d = ds["dataset_id"]
    a, b = _tag(store, d, "Pacing", definition="Speed"), _tag(store, d, "Recommends")
    call(store, "tags.apply", dataset_id=d, row_id=0, variable="Q10", tag_ids=[a["id"], b["id"]])
    call(store, "tags.apply", dataset_id=d, row_id=1, variable="Q10", tag_ids=[b["id"]])
    n = len(store.get(d).df)
    x = call(store, "export.qualitative", dataset_id=d, variable="Q10", kind="responses", format="xlsx",
             path=str(tmp_path / "coded.xlsx"), context_variables=["Q6"])
    assert x["n_responses"] == n and x["bytes"] == (tmp_path / "coded.xlsx").stat().st_size
    wb = load_workbook(tmp_path / "coded.xlsx")
    assert wb.sheetnames == ["Coded responses", "Codebook", "Summary"]
    rows = list(wb["Coded responses"].values)
    assert rows[0] == ("Row ID", "Q6", "Response", "Tags", "Pacing", "Recommends")
    assert rows[1][3:] == ("Pacing, Recommends", 1, 1) and rows[2][3:] == ("Recommends", 0, 1)
    assert len(rows) == n + 1
    call(store, "export.qualitative", dataset_id=d, variable="Q10", kind="responses", format="docx",
         path=str(tmp_path / "coded.docx"))
    heads = [p.text for p in Document(str(tmp_path / "coded.docx")).paragraphs if p.style.name == "Heading 1"]
    assert heads == ["Pacing (1 responses, 0.9%)", "Recommends (2 responses, 1.8%)", f"Not tagged yet ({n - 2} responses)"]
    call(store, "export.qualitative", dataset_id=d, variable="Q10", kind="codebook", format="xlsx",
         path=str(tmp_path / "book.xlsx"))
    assert list(load_workbook(tmp_path / "book.xlsx")["Tag codebook"].values)[1] == ("Pacing", "Speed", "#0072B2", 1, "0.9%")
    call(store, "export.qualitative", dataset_id=d, variable="Q10", kind="codebook", format="docx",
         path=str(tmp_path / "book.docx"))
    assert Document(str(tmp_path / "book.docx")).tables[0].rows[2].cells[0].text == "Recommends"
    with pytest.raises(InvalidParams):
        call(store, "export.qualitative", dataset_id=d, variable="Q10", kind="responses", format="xlsx",
             path=str(tmp_path / "coded.xlsx"))  # exists, no overwrite


def test_over_stdio(engine, tmp_path):
    """The whole flow through the real JSON-RPC server."""
    pv = engine.call("dataset.import_preview", {"files": [{"path": str(MESSY / "messy_3header.csv"), "sheet_name": None}],
                                                "qualtrics_mode": "auto", "stack_onto_dataset_id": None})
    f = pv["files"][0]
    meta = engine.call("dataset.import", {
        "preview_id": pv["preview_id"], "row_filters": [], "variables": [], "stack": None,
        "files": [{"file_id": f["file_id"], "sheet_name": None, "encoding": f["encoding"], "delimiter": f["delimiter"],
                   "qualtrics_header_rows": f["qualtrics"]["header_rows"], "time_label": None, "drop_columns": []}],
    })["dataset_meta"]
    d = meta["dataset_id"]
    tag = engine.call("tags.codebook.upsert", {"dataset_id": d, "tag": {"name": "Pacing"}})["tag"]
    engine.call("tags.apply", {"dataset_id": d, "row_id": 2, "variable": "Q10", "tag_ids": [tag["id"]]})
    s = engine.call("tags.summary", {"dataset_id": d, "variable": "Q10", "by": "Q6"})
    assert s["overall"][0]["count"] == 1
    err = engine.error("tags.responses", {"dataset_id": d, "variable": "Q10", "limit": 0})
    assert err["code"] == -32003
