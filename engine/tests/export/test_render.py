"""Golden tests for Phase 8 exports (SPEC §10.3): open the produced DOCX (python-docx), XLSX (openpyxl),
CSV and PDF (pypdf) and assert structure and content. Handlers are called in-process with a
DatasetStore; test_rpc_export.py covers the same methods through the real stdio server.

HTML snapshot: regenerate with
    STATLY_UPDATE_SNAPSHOTS=1 .venv/bin/python -m pytest tests/export/test_render.py
and review the diff of tests/export/snapshots/*.html.
"""

from __future__ import annotations

import base64
import csv
import io
import json
import os
from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from statly_engine.data.store import DatasetStore
from statly_engine.errors import InvalidParams
from statly_engine.export import plain
from statly_engine.rpc_methods import export as ex

HERE = Path(__file__).parent
REPO = HERE.parents[2]
EXAMPLE = json.loads((REPO / "contracts" / "examples" / "AnalysisResult.json").read_text())
REAL = json.loads((HERE / "data" / "t_test_independent_q3_10.json").read_text())  # real t_test.independent
META = json.loads((REPO / "contracts" / "examples" / "DatasetMeta.json").read_text())
LOG = json.loads((REPO / "contracts" / "examples" / "TestLogEntry.json").read_text())
SNAP = HERE / "snapshots"


def _png_b64(w=600, h=400) -> str:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (200, 220, 240)).save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode()


def _italic_texts(paragraphs) -> set[str]:
    return {r.text for p in paragraphs for r in p.runs if r.italic}


def _norm(s: str) -> str:
    return " ".join(s.split())


@pytest.fixture
def store():
    return DatasetStore()


# --- HTML ---------------------------------------------------------------------------------------
def test_table_html_snapshot(store):
    out = ex.table_html(store, {"apa_table": REAL["apa_table"]})
    snap = SNAP / "table_t_test_independent.html"
    if os.environ.get("STATLY_UPDATE_SNAPSHOTS") or not snap.exists():
        SNAP.mkdir(exist_ok=True)
        snap.write_text(out["html"], encoding="utf-8")
    assert out["html"] == snap.read_text(encoding="utf-8")
    html = out["html"]
    assert "Times New Roman" in html and "12pt" in html
    assert "<b>Table 1</b>" in html and "<i>Comparison of Q3_10 by Time point</i>" in html
    assert "&lt; .001" in html and "<i>SD</i>" in html and "<i>Note.</i>" in html
    assert "border-left" not in html and "border-right" not in html  # horizontal rules only
    lines = out["plain_text"].splitlines()
    assert lines[:2] == ["Table 1", "Comparison of Q3_10 by Time point"]
    assert lines[3].split("\t")[:4] == ["Variable", "n", "M", "SD"]
    assert "< .001" in lines[4].split("\t")


def test_table_html_number_override_and_validation(store):
    out = ex.table_html(store, {"apa_table": EXAMPLE["apa_table"], "number": 7})
    assert "<b>Table 7</b>" in out["html"] and out["plain_text"].startswith("Table 7\n")
    with pytest.raises(InvalidParams):
        ex.table_html(store, {"apa_table": {"title": "x"}})


# --- Report -------------------------------------------------------------------------------------
def _report(store, tmp_path, fmt, **kw):
    params = {"title": "Attitude Study", "author": "A. Student", "results": [EXAMPLE, REAL], "format": fmt,
              "path": str(tmp_path / f"report.{fmt}"),
              "charts": [{"request_id": REAL["inputs"]["request"]["request_id"], "png_base64": _png_b64(),
                          "title": "Q3_10 by Time Point"}], **kw}
    return ex.report(store, params)


def test_report_docx_golden(store, tmp_path):
    out = _report(store, tmp_path, "docx")
    assert out["bytes"] == Path(out["path"]).stat().st_size > 0
    doc = Document(out["path"])
    n_tables = sum((1 if r["apa_table"] else 0) + len(r["additional_tables"]) for r in (EXAMPLE, REAL))
    assert len(doc.tables) == n_tables
    assert all(t.style.name == "APA Table" for t in doc.tables)
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Attitude Study"
    assert [t for t in texts if t.startswith("Table ")] == [f"Table {i}" for i in range(1, n_tables + 1)]
    assert "Figure 1" in texts and "Q3_10 by Time Point" in texts and len(doc.inline_shapes) == 1
    # APA sentences verbatim, symbols italic
    for res in (EXAMPLE, REAL):
        assert plain(res["apa_sentence"]) in texts
        assert res["plain_language_summary"] in texts
    assert {"t", "p", "g", "M", "SD"} <= _italic_texts(doc.paragraphs)
    # the real result's table: two header rows (spanning Pre/Post), italic symbols, p string preserved
    real = doc.tables[-1]
    assert [c.text for c in real.rows[0].cells][1:4] == ["Pre"] * 3  # merged span
    hdr = real.rows[1]
    assert [c.text for c in hdr.cells][:4] == ["Variable", "n", "M", "SD"]
    assert all(r.italic for c in hdr.cells[1:4] for p in c.paragraphs for r in p.runs)
    body = [c.text for c in real.rows[2].cells]
    assert "< .001" in body and "[-1.00, -0.27]" in body and "117.28" in body
    # notes: 'Note.' italic then the engine's general note
    note = next(p for p in doc.paragraphs if p.text.startswith("Note. Welch's t test"))
    assert note.runs[0].text == "Note." and note.runs[0].italic
    assert "Assumption Checks" in texts
    assert any("Shapiro-Wilk" in t and "p < .001" in t and "Not met" in t for t in texts)
    assert doc.core_properties.author == "A. Student"


def test_report_include_flags(store, tmp_path):
    out = _report(store, tmp_path, "docx", include={"tables": False, "assumptions": False, "charts": False,
                                                    "sentences": True}, charts=None)
    doc = Document(out["path"])
    texts = [p.text for p in doc.paragraphs]
    assert len(doc.tables) == 0 and len(doc.inline_shapes) == 0 and "Assumption Checks" not in texts
    assert plain(REAL["apa_sentence"]) in texts


def test_report_pdf_smoke(store, tmp_path):
    out = _report(store, tmp_path, "pdf")
    reader = PdfReader(out["path"])
    assert len(reader.pages) > 0
    text = _norm(" ".join(p.extract_text() for p in reader.pages))
    assert _norm(plain(REAL["apa_sentence"])) in text
    assert "Table 3" in text and "Figure 1" in text and "< .001" in text and "Attitude Study" in text
    assert reader.metadata.title == "Attitude Study"


def test_report_rejects_bad_input(store, tmp_path):
    with pytest.raises(InvalidParams):
        _report(store, tmp_path, "docx", results=[], charts=[])  # nothing to put in the document
    with pytest.raises(InvalidParams):
        _report(store, tmp_path, "docx", charts=[{"request_id": "x", "png_base64": base64.b64encode(b"GIF89a").decode()}])
    with pytest.raises(InvalidParams):
        _report(store, tmp_path, "docx", include={"tables": "yes"})
    with pytest.raises(InvalidParams):
        _report(store, tmp_path, "rtf")
    bad = json.loads(json.dumps(REAL))
    del bad["apa_sentence"]
    with pytest.raises(InvalidParams):
        _report(store, tmp_path, "docx", results=[bad])


# --- Paths --------------------------------------------------------------------------------------
def test_path_rules(store, tmp_path):
    base = {"title": "R", "results": [EXAMPLE], "format": "docx"}
    for path, reason in (("report.docx", "relative_path"), (str(tmp_path / "r.pdf"), "wrong_extension"),
                         (str(tmp_path / "nope" / "r.docx"), "missing_folder"),
                         (str(tmp_path / "a" / ".." / "r.docx"), "relative_path")):
        with pytest.raises(InvalidParams) as err:
            ex.report(store, {**base, "path": path})
        assert err.value.data["reason"] == reason
    target = tmp_path / "r.docx"
    target.write_bytes(b"keep me")
    with pytest.raises(InvalidParams) as err:
        ex.report(store, {**base, "path": str(target)})
    assert err.value.data["reason"] == "file_exists" and target.read_bytes() == b"keep me"
    out = ex.report(store, {**base, "path": str(target), "overwrite": True})
    assert out["bytes"] > 7 and Document(str(target)).tables
    assert not (tmp_path / "r.docx.tmp").exists()


# --- Dataset + codebook -------------------------------------------------------------------------
def _store_with_data() -> tuple[DatasetStore, str]:
    store = DatasetStore()
    df = pd.DataFrame({
        "_statly_row_id": pd.array([0, 1, 2], dtype="int64"),
        "ResponseId": pd.array(["R_1", "R_2", "=cmd()"], dtype="string"),
        "IPAddress": pd.array(["1.1.1.1", "2.2.2.2", "3.3.3.3"], dtype="string"),
        "student_code": pd.array(["ab12", "cd34", None], dtype="string"),
        "Time": pd.array(["Pre", "Post", "Pre"], dtype="string"),
        "Q5_1": pd.array([4, 5, None], dtype="Int64"),
        "Q5_2": pd.array([2, -99, 1], dtype="Int64"),
        "math_attitude": [4.0, float("nan"), 5.0],
        "score_gain": [1.5, 2.0, float("nan")],
        "SC0_band": pd.array(["Low", "High", None], dtype="string"),
    })
    meta = {**META, "variables": [{**v, "is_metadata": v["name"] == "IPAddress"} for v in META["variables"]]}
    store.commit(META["dataset_id"], df, meta, {})
    return store, META["dataset_id"]


def test_data_xlsx_golden(tmp_path):
    store, ds = _store_with_data()
    out = ex.data(store, {"dataset_id": ds, "format": "xlsx", "path": str(tmp_path / "d.xlsx"),
                          "include_metadata_columns": False, "options": {"label_row": True}})
    assert out["n_rows"] == 3 and out["n_columns"] == len(META["variables"]) - 1
    ws = load_workbook(out["path"]).active
    rows = list(ws.iter_rows(values_only=True))
    assert "IPAddress" not in rows[0] and "_statly_row_id" not in rows[0]
    assert rows[0][:3] == ("ResponseId", "student_code", "Time")
    assert rows[1][:3] == ("Response ID", "Self-generated ID", "Time point")  # label row
    assert ws["A1"].font.b and ws["A2"].font.i
    q52 = rows[0].index("Q5_2")
    assert [r[q52] for r in rows[2:]] == [2, None, 1]  # -99 declared missing -> blank
    assert rows[4][0] == "=cmd()" and ws.cell(row=5, column=1).data_type == "s"  # text, not a formula
    assert rows[3][rows[0].index("math_attitude")] is None
    assert out["pii_columns"] == [v["name"] for v in META["variables"] if v["is_pii"] and v["name"] != "IPAddress"]


def test_data_csv_options(tmp_path):
    store, ds = _store_with_data()
    out = ex.data(store, {"dataset_id": ds, "format": "csv", "path": str(tmp_path / "d.csv"),
                          "include_metadata_columns": True,
                          "options": {"blank_missing_codes": False, "exclude_pii": True}})
    raw = Path(out["path"]).read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    pii = {v["name"] for v in META["variables"] if v["is_pii"]}
    assert pii and not pii & set(rows[0])
    q52 = rows[0].index("Q5_2")
    assert [r[q52] for r in rows[1:]] == ["2", "-99", "1"]
    assert rows[1][rows[0].index("math_attitude")] == "4" and rows[2][rows[0].index("math_attitude")] == ""
    assert out["pii_columns"] == []


def test_codebook_xlsx_and_docx(tmp_path):
    store, ds = _store_with_data()
    out = ex.codebook(store, {"dataset_id": ds, "format": "xlsx", "path": str(tmp_path / "cb.xlsx")})
    wb = load_workbook(out["path"])
    assert wb.sheetnames == ["Codebook", "Scales"]
    rows = list(wb["Codebook"].iter_rows(values_only=True))
    assert rows[0][:6] == ("Name", "Label", "Question text", "Role", "Level", "Value labels")
    by = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert len(by) == len(META["variables"])
    assert by["Q5_2"]["Reverse-scored"] == "Yes (1–5, scored as 6 − x)" and by["Q5_2"]["Missing codes"] == "-99"
    assert by["Q5_1"]["Value labels"].startswith("1 = Strongly disagree; 2 = Disagree")
    assert by["Q5_1"]["Scale"] == "Math attitude" and by["Q5_1"]["Scoring rule"].startswith("Mean of 2 items")
    assert by["math_attitude"]["Computed as"].startswith("Mean of 2 items (Q5_1, Q5_2)")
    assert by["score_gain"]["Computed as"] == "Difference score: SC0 at Post minus SC0 at Pre."
    assert by["SC0_band"]["Computed as"].startswith("Recode of SC0: 0 to 5 → Low")
    assert list(wb["Scales"].iter_rows(values_only=True))[1] == ("Math attitude", "Q5_1, Q5_2", "Mean", "2",
                                                                  "math_attitude")
    out = ex.codebook(store, {"dataset_id": ds, "format": "docx", "path": str(tmp_path / "cb.docx")})
    doc = Document(out["path"])
    assert len(doc.tables) == 2 and doc.tables[0].rows[0].cells[0].text == "Name"
    assert len(doc.tables[0].rows) == 1 + len(META["variables"])


# --- Test Log -----------------------------------------------------------------------------------
def test_test_log_all_formats(store, tmp_path):
    second = json.loads(json.dumps(LOG))
    second["result_summary"]["p"] = 0.0004
    second["correction_method"], second["adjusted_p"], second["family_id"] = "none", None, None
    entries = [LOG, second]
    x = ex.test_log(store, {"entries": entries, "format": "xlsx", "path": str(tmp_path / "log.xlsx")})
    rows = list(load_workbook(x["path"]).active.iter_rows(values_only=True))
    h = rows[0]
    assert h[:7] == ("#", "Date", "Analysis", "Outcome(s)", "Statistic", "p", "Effect size")
    r1, r2 = dict(zip(h, rows[1])), dict(zip(h, rows[2]))
    assert r1["Statistic"] == "t(37.40) = 2.31" and r1["p"] == ".026" and r1["Adjusted p"] == ".052"
    assert r1["Effect size"] == "g = 0.72, 95% CI [0.08, 1.35]" and r1["Correction"] == "Holm"
    assert r2["p"] == "< .001" and r2["Adjusted p"] is None and r2["Correction"] == "None"
    assert r1["APA sentence"] == plain(LOG["result_summary"]["apa_sentence"])
    c = ex.test_log(store, {"entries": entries, "format": "csv", "path": str(tmp_path / "log.csv")})
    crow = list(csv.reader(io.StringIO(Path(c["path"]).read_text(encoding="utf-8-sig"))))
    assert crow[0][0] == "#" and crow[2][5] == "< .001"
    d = ex.test_log(store, {"entries": entries, "format": "docx", "path": str(tmp_path / "log.docx")})
    doc = Document(d["path"])
    t = doc.tables[0]
    assert len(t.rows) == 3 and t.rows[2].cells[5].text == "< .001"
    stat_runs = t.rows[1].cells[4].paragraphs[0].runs
    assert stat_runs[0].text == "t" and stat_runs[0].italic
    with pytest.raises(InvalidParams):
        ex.test_log(store, {"entries": [{"id": "x"}], "format": "csv", "path": str(tmp_path / "bad.csv")})
