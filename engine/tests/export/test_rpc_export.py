"""export.* through the real stdio server: practice data -> t_test.independent -> every export format."""

from __future__ import annotations

import base64
import io
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

REPO = Path(__file__).resolve().parents[3]
PRACTICE = REPO / "fixtures" / "practice" / "one_group_prepost_likert"
LOG = json.loads((REPO / "contracts" / "examples" / "TestLogEntry.json").read_text())


def _import(engine) -> dict:
    files = [PRACTICE / "pre.csv", PRACTICE / "post.csv"]
    pv = engine.call("dataset.import_preview", {"files": [{"path": str(p), "sheet_name": None} for p in files],
                                                "qualtrics_mode": "auto", "stack_onto_dataset_id": None})
    labels = ["Pre", "Post"]
    decisions = [{"file_id": f["file_id"], "sheet_name": f["sheet_name"], "encoding": f["encoding"],
                  "delimiter": f["delimiter"], "qualtrics_header_rows": f["qualtrics"]["header_rows"],
                  "time_label": lbl, "drop_columns": []} for f, lbl in zip(pv["files"], labels)]
    return engine.call("dataset.import", {
        "preview_id": pv["preview_id"], "files": decisions, "row_filters": [], "variables": [],
        "stack": {"time_variable": "Time", "column_matches": pv["stack_proposal"],
                  "levels": [{"file_id": f["file_id"], "label": lbl} for f, lbl in zip(pv["files"], labels)]},
    })["dataset_meta"]


def _run(engine, meta: dict) -> dict:
    return engine.call("analysis.run", {
        "schema_version": 1, "request_id": "req-export", "analysis_id": "t_test.independent",
        "dataset_id": meta["dataset_id"], "snapshot_id": meta["snapshot_id"],
        "variables": {"outcome": ["Q3_10"], "group": ["Time"]}, "subset": [], "options": {}, "corrections": [],
        "alpha": 0.05, "tails": "two_sided", "ci_level": 0.95})


def test_all_exports_over_rpc(engine, tmp_path):
    meta = _import(engine)
    res = _run(engine, meta)

    html = engine.call("export.table_html", {"apa_table": res["apa_table"]})
    assert "&lt; .001" in html["html"] and "< .001" in html["plain_text"]

    buf = io.BytesIO()
    Image.new("RGB", (800, 500), "white").save(buf, "PNG")
    chart = {"request_id": "req-export", "png_base64": base64.b64encode(buf.getvalue()).decode()}
    for fmt in ("docx", "pdf"):
        out = engine.call("export.report", {"title": "Practice Report", "author": None, "results": [res],
                                            "include": {"charts": True}, "charts": [chart], "format": fmt,
                                            "path": str(tmp_path / f"report.{fmt}")})
        assert out["bytes"] == Path(out["path"]).stat().st_size
    doc = Document(str(tmp_path / "report.docx"))
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1
    assert "< .001" in [c.text for c in doc.tables[0].rows[2].cells]
    sentence = "".join(r["text"] for r in res["apa_sentence"])
    pdf_text = " ".join(" ".join(p.extract_text() for p in PdfReader(str(tmp_path / "report.pdf")).pages).split())
    assert " ".join(sentence.split()) in pdf_text

    for fmt in ("xlsx", "csv"):
        out = engine.call("export.data", {"dataset_id": meta["dataset_id"], "format": fmt,
                                          "path": str(tmp_path / f"data.{fmt}"), "include_metadata_columns": False,
                                          "options": {"label_row": True}})
        assert out["n_rows"] == meta["n_rows"] == 120 and out["snapshot_id"] == meta["snapshot_id"]
    rows = list(load_workbook(tmp_path / "data.xlsx").active.iter_rows(values_only=True))
    assert len(rows) == 2 + 120 and "Q3_10" in rows[0] and "Time" in rows[0]
    assert not any(v["is_metadata"] and v["name"] in rows[0] for v in meta["variables"])

    for fmt in ("xlsx", "docx"):
        out = engine.call("export.codebook", {"dataset_id": meta["dataset_id"], "format": fmt,
                                              "path": str(tmp_path / f"codebook.{fmt}")})
        assert out["bytes"] > 0
    cb = list(load_workbook(tmp_path / "codebook.xlsx")["Codebook"].iter_rows(values_only=True))
    assert len(cb) == 1 + len(meta["variables"])

    for fmt in ("xlsx", "csv", "docx"):
        out = engine.call("export.test_log", {"entries": [LOG], "format": fmt, "path": str(tmp_path / f"log.{fmt}")})
        assert out["bytes"] > 0

    # never overwrite silently; refuse relative paths; unknown dataset
    err = engine.error("export.test_log", {"entries": [LOG], "format": "csv", "path": str(tmp_path / "log.csv")})
    assert err["code"] == -32003 and err["data"]["reason"] == "file_exists"
    ok = engine.call("export.test_log", {"entries": [LOG], "format": "csv", "path": str(tmp_path / "log.csv"),
                                         "overwrite": True})
    assert ok["bytes"] > 0
    err = engine.error("export.report", {"title": "R", "results": [res], "format": "pdf", "path": "r.pdf"})
    assert err["code"] == -32003 and err["data"]["reason"] == "relative_path"
    err = engine.error("export.data", {"dataset_id": "nope", "format": "csv", "path": str(tmp_path / "x.csv")})
    assert err["code"] == -32002
    assert sorted(p.name for p in tmp_path.iterdir()) == sorted(
        ["report.docx", "report.pdf", "data.xlsx", "data.csv", "codebook.xlsx", "codebook.docx",
         "log.xlsx", "log.csv", "log.docx"])


def test_report_shows_holm_adjusted_p_for_family_member(engine, tmp_path):
    """A DOCX report for two logged tests, one in a Holm family, shows the adjusted p (SPEC §9/§10.3)."""
    meta = _import(engine)
    res_a = _run(engine, meta)
    res_b = engine.call("analysis.run", {
        "schema_version": 1, "request_id": "req-export-2", "analysis_id": "t_test.independent",
        "dataset_id": meta["dataset_id"], "snapshot_id": meta["snapshot_id"],
        "variables": {"outcome": ["Q3_9"], "group": ["Time"]}, "subset": [], "options": {}, "corrections": [],
        "alpha": 0.05, "tails": "two_sided", "ci_level": 0.95})

    def entry(res, req_id, family_id, adjusted_p):
        return {
            "schema_version": 1, "id": req_id, "timestamp": "2026-09-24T21:00:00Z",
            "request": res["inputs"]["request"],
            "result_summary": {
                "analysis_label": "Independent-samples t test",
                "outcome_variables": list(res["inputs"]["request"]["variables"]["outcome"]),
                "primary_statistic": res["statistics"][0] if res["statistics"] else None,
                "p": res["statistics"][0]["p"] if res["statistics"] else None,
                "primary_effect_size": res["effect_sizes"][0] if res["effect_sizes"] else None,
                "n_used": res["inputs"]["n_used"], "apa_sentence": res["apa_sentence"],
                "plain_language_summary": res["plain_language_summary"], "engine_version": res["engine_version"],
            },
            "result_path": None, "family_id": family_id,
            "correction_method": "holm" if family_id else "none", "adjusted_p": adjusted_p,
        }

    log = [entry(res_a, "req-export", "fam-1", 0.048), entry(res_b, "req-export-2", None, None)]
    families = [{"id": "fam-1", "name": "Attitude items"}]
    out = engine.call("export.report", {
        "title": "Family Report", "results": [res_a, res_b], "include": {"sentences": True},
        "test_log": log, "test_families": families, "format": "docx",
        "path": str(tmp_path / "family_report.docx")})
    assert out["bytes"] > 0
    doc = Document(str(tmp_path / "family_report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Holm-adjusted p = .048 (family: Attitude items)." in text
    # the second (non-family) result must not get an adjusted-p note
    assert text.count("Holm-adjusted") == 1


def test_chart_only_report_for_chart_builder(engine, tmp_path):
    """Chart Builder's "Save figure > PDF" sends charts with no analyses (results: [])."""
    buf = io.BytesIO()
    Image.new("RGB", (600, 400), "white").save(buf, "PNG")
    chart = {"request_id": "chart-builder", "png_base64": base64.b64encode(buf.getvalue()).decode(),
             "title": "Math attitude by score band"}
    include = {"tables": False, "sentences": False, "assumptions": False, "charts": True}
    for fmt in ("pdf", "docx"):
        out = engine.call("export.report", {"title": "Math attitude by score band", "results": [],
                                            "include": include, "charts": [chart], "format": fmt,
                                            "path": str(tmp_path / f"figure.{fmt}")})
        assert out["bytes"] == Path(out["path"]).stat().st_size
    pages = PdfReader(str(tmp_path / "figure.pdf")).pages
    assert sum(len(p.images) for p in pages) == 1
    assert "Figure 1" in " ".join(p.extract_text() for p in pages)
    assert len(Document(str(tmp_path / "figure.docx")).inline_shapes) == 1

    # Still an error when there is nothing to put in the document.
    err = engine.error("export.report", {"title": "Empty", "results": [], "include": include, "charts": [],
                                         "format": "pdf", "path": str(tmp_path / "empty.pdf")})
    assert "at least one analysis" in err["message"]
    err = engine.error("export.report", {"title": "Empty", "results": [], "include": {"charts": False},
                                         "charts": [chart], "format": "pdf", "path": str(tmp_path / "e2.pdf")})
    assert "at least one analysis" in err["message"]
