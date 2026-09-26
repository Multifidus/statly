"""export.plan (Study Planner DOCX, SPEC §11.2): structure of the document and param validation."""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from docx import Document

from statly_engine.errors import InvalidParams
from statly_engine.rpc_methods import export as export_rpc

REPO = Path(__file__).resolve().parents[3]
PLAN = json.loads((REPO / "contracts" / "examples" / "StudyPlan.json").read_text())


def _plan() -> dict:
    p = copy.deepcopy(PLAN)
    p["design"]["answers"]["planner_research_question"] = "Does the intervention improve math attitude?"
    p["recommendations"].append({"category": "qualtrics_setup", "text": "Check the recode values.",
                                 "why": "Scores must be 1-5 in the order you expect."})
    p["power_analyses"].append({
        "analysis_id": "t_test.independent", "mode": "sensitivity",
        "inputs": {"alpha": 0.05, "power": 0.8, "tails": "two_sided", "effect_size_metric": "d",
                   "effect_size": None, "n_total": 60, "allocation_ratio": 1},
        "outputs": {"n_total": 60, "n_per_group": [30, 30], "detectable_effect": 0.7338, "achieved_power": None,
                    "method_note": "Noncentral t."}})
    return p


def _text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_plan_docx_structure(tmp_path):
    out = tmp_path / "plan.docx"
    res = export_rpc.plan(None, {
        "plan": _plan(), "format": "docx", "path": str(out),
        "labels": {"ancova": "ANCOVA", "quade": "Quade test", "normality": "Normality",
                   "partial_eta_sq": "Partial eta squared"},
        "interview": [{"question": "What do you want to know?", "answer": "Did scores differ between groups?"}]})
    assert res["path"] == str(out) and res["bytes"] == out.stat().st_size > 0
    doc = Document(str(out))
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.paragraphs[0].text == "Math attitude intervention study"
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1 == ["Research question", "Design summary", "Planned analyses", "Sample size (power analysis)",
                  "Data collection recommendations", "Assumptions to check after you collect data"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert "ANCOVA (posttest adjusted for pretest)" in h2
    assert {"Sample size needed", "What your sample can detect", "Collecting your data",
            "Setting up your Qualtrics survey"} <= set(h2)
    text = _text(doc)
    assert "Does the intervention improve math attitude?" in text
    assert "Quade test" in text and "Partial eta squared" in text
    assert "plan for about 128 people in total" in text
    assert "d = 0.73" in text
    assert "Total sample needed (N)" in text and "64, 64" in text
    assert "Did scores differ between groups?" in text
    assert "☐ Normality" in text and "☐ Homogeneity of regression slopes" in text
    # interview grid + one grid per power analysis
    assert len(doc.tables) == 3


def test_plan_without_optional_parts(tmp_path):
    p = copy.deepcopy(PLAN)
    p["planned_analyses"], p["power_analyses"], p["recommendations"] = [], [], []
    out = tmp_path / "empty.docx"
    export_rpc.plan(None, {"plan": p, "format": "docx", "path": str(out)})
    text = _text(Document(str(out)))
    assert "No analyses planned yet." in text and "No power analysis yet." in text
    assert "Research question" not in text


@pytest.mark.parametrize("patch, msg", [
    ({"format": "pdf"}, "format"),
    ({"plan": {"title": "x"}}, "contract"),
    ({"labels": {"a": 1}}, "labels"),
    ({"interview": [{"question": "q"}]}, "interview"),
])
def test_plan_invalid_params(tmp_path, patch, msg):
    params = {"plan": _plan(), "format": "docx", "path": str(tmp_path / "x.docx")} | patch
    with pytest.raises(InvalidParams) as e:
        export_rpc.plan(None, params)
    assert msg in str(e.value)


def test_plan_refuses_existing_file_without_overwrite(tmp_path):
    out = tmp_path / "plan.docx"
    out.write_bytes(b"old")
    with pytest.raises(InvalidParams):
        export_rpc.plan(None, {"plan": _plan(), "format": "docx", "path": str(out)})
    export_rpc.plan(None, {"plan": _plan(), "format": "docx", "path": str(out), "overwrite": True})
    assert out.stat().st_size > 3


def test_plan_via_rpc_server(engine, tmp_path):
    out = tmp_path / "srv.docx"
    res = engine.call("export.plan", {"plan": _plan(), "format": "docx", "path": str(out)})
    assert res["bytes"] > 0 and out.exists()
